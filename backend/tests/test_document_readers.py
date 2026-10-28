from __future__ import annotations

import hashlib
import io
import os
import subprocess
from contextlib import closing
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

import openpyxl
import pymupdf
import pytest
import xlwt
from PIL import Image, ImageDraw, ImageFont

from app.services import ocr as ocr_module
from app.services.document_reader import read_document
from app.services.ocr import TesseractOCR
from app.services.reading_types import DocumentSource, ReadingSettings
from app.workers.jobs import ProcessingError


def source_for(path: Path) -> DocumentSource:
    return DocumentSource(
        document_id=uuid4(),
        original_filename=path.name,
        storage_key=path.name,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        extension=path.suffix.lower(),
    )


def options(**overrides) -> ReadingSettings:
    return ReadingSettings(_env_file=None, **overrides)


def invoice_image() -> Image.Image:
    image = Image.new("RGB", (1200, 350), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=60)
    draw.text((40, 40), "INVOICE 12345", fill="black", font=font)
    draw.text((40, 150), "TOTAL 150", fill="black", font=font)
    return image


def save_scanned_pdf(path: Path) -> None:
    with invoice_image() as image:
        output = io.BytesIO()
        image.save(output, format="PNG")

    with pymupdf.open() as document:
        page = document.new_page(width=600, height=175)
        page.insert_image(page.rect, stream=output.getvalue())
        document.save(str(path))


def test_native_pdf_streams_1000_pages_without_rendering(tmp_path, monkeypatch):
    path = tmp_path / "large-native.pdf"

    with pymupdf.open() as document:
        for number in range(1, 1001):
            page = document.new_page()
            page.insert_text(
                (40, 50),
                f"SYNTHETIC DIGITAL INVOICE PAGE {number} TOTAL 125",
            )
        document.save(str(path))

    calls = []
    original_get_text = pymupdf.Page.get_text

    def observed_get_text(page, *args, **kwargs):
        calls.append(page.number)
        return original_get_text(page, *args, **kwargs)

    def forbidden_render(*args, **kwargs):
        raise AssertionError("Native PDF pages should not be rendered.")

    def forbidden_ocr(image):
        raise AssertionError("Native PDF pages should not require OCR.")

    monkeypatch.setattr(pymupdf.Page, "get_text", observed_get_text)
    monkeypatch.setattr(pymupdf.Page, "get_pixmap", forbidden_render)

    source = source_for(path)

    with closing(
        read_document(path, source, settings=options(), ocr=forbidden_ocr)
    ) as units:
        first = next(units)
        assert calls == [0]
        assert first.position.page_number == 1
        assert first.position.part_count == 1000
        assert first.source.document_id == source.document_id
        assert "TOTAL 125" in first.text

        remaining_count = sum(1 for _ in units)

    assert remaining_count == 999
    assert len(calls) == 1000


def test_scanned_pdf_routes_to_injected_ocr(tmp_path):
    path = tmp_path / "scan.pdf"
    save_scanned_pdf(path)
    calls = []

    def mock_ocr(image):
        calls.append(image.size)
        return "MOCK OCR ROUTING RESULT"

    units = list(
        read_document(path, source_for(path), settings=options(), ocr=mock_ocr)
    )

    assert len(calls) == 1
    assert units[0].method == "ocr"
    assert units[0].native_text == ""
    assert units[0].ocr_text == "MOCK OCR ROUTING RESULT"
    assert "ocr_output_requires_review" in units[0].warnings


def test_mixed_pdf_preserves_native_and_ocr_evidence(tmp_path):
    path = tmp_path / "mixed.pdf"

    with invoice_image() as image:
        output = io.BytesIO()
        image.save(output, format="PNG")

    with pymupdf.open() as document:
        page = document.new_page(width=600, height=500)
        page.insert_text(
            (20, 30),
            "SELECTABLE SYNTHETIC HEADER REFERENCE",
        )
        page.insert_image(
            pymupdf.Rect(20, 80, 580, 300),
            stream=output.getvalue(),
        )
        document.save(str(path))

    units = list(
        read_document(
            path,
            source_for(path),
            settings=options(),
            ocr=lambda image: "MOCK IMAGE REGION TEXT",
        )
    )

    assert units[0].method == "pdf_text_and_ocr"
    assert "SELECTABLE SYNTHETIC HEADER" in units[0].native_text
    assert units[0].ocr_text == "MOCK IMAGE REGION TEXT"
    assert "native_and_ocr_may_overlap" in units[0].warnings


@pytest.mark.parametrize(
    ("extension", "image_format"),
    [(".jpg", "JPEG"), (".png", "PNG"), (".tiff", "TIFF")],
)
def test_image_formats_route_to_ocr(tmp_path, extension, image_format):
    path = tmp_path / f"invoice{extension}"

    with invoice_image() as image:
        image.save(path, format=image_format)

    units = list(
        read_document(
            path,
            source_for(path),
            settings=options(),
            ocr=lambda image: "MOCK IMAGE OCR",
        )
    )

    assert len(units) == 1
    assert units[0].position.frame_number == 1
    assert units[0].method == "ocr"


def test_multipage_tiff_preserves_frame_numbers(tmp_path):
    path = tmp_path / "multipage.tiff"

    with invoice_image() as first:
        with first.copy() as second:
            first.save(
                path,
                format="TIFF",
                save_all=True,
                append_images=[second],
            )

    calls = []

    def mock_ocr(image):
        calls.append(image.size)
        return f"MOCK FRAME {len(calls)}"

    units = list(
        read_document(path, source_for(path), settings=options(), ocr=mock_ocr)
    )

    assert len(calls) == 2
    assert [unit.position.frame_number for unit in units] == [1, 2]
    assert [unit.sequence for unit in units] == [1, 2]


def test_empty_ocr_output_is_not_invented(tmp_path):
    path = tmp_path / "blank.png"

    with Image.new("RGB", (100, 100), "white") as image:
        image.save(path)

    unit = next(
        read_document(
            path,
            source_for(path),
            settings=options(),
            ocr=lambda image: "",
        )
    )

    assert unit.text == ""
    assert "no_text_detected" in unit.warnings


def test_xlsx_preserves_sheets_rows_numbers_and_formulas(tmp_path):
    path = tmp_path / "invoice.xlsx"
    workbook = openpyxl.Workbook()

    try:
        sheet = workbook.active
        sheet.title = "Invoices"
        sheet.append(["invoice", "total"])
        sheet.append(["SYN-001", 125])
        sheet.append(["calculation", "=SUM(B2:B2)"])
        workbook.create_sheet("Empty Sheet")
        workbook.save(path)
    finally:
        workbook.close()

    units = list(read_document(path, source_for(path), settings=options()))

    sheet_units = [unit for unit in units if unit.kind == "sheet"]
    assert [unit.position.sheet_name for unit in sheet_units] == [
        "Invoices",
        "Empty Sheet",
    ]

    row_two = next(
        unit
        for unit in units
        if unit.kind == "row" and unit.position.row_number == 2
    )
    assert row_two.cells[1].address == "B2"
    assert row_two.cells[1].data_type == "number"
    assert Decimal(row_two.cells[1].value) == Decimal("125")

    formula_row = next(
        unit
        for unit in units
        if unit.kind == "row" and unit.position.row_number == 3
    )
    assert formula_row.cells[1].value is None
    assert formula_row.cells[1].formula == "=SUM(B2:B2)"
    assert "formula_value_unavailable" in formula_row.warnings


def test_legacy_xls_reads_multiple_sheets(tmp_path):
    path = tmp_path / "invoice.xls"
    workbook = xlwt.Workbook()

    for name in ("Invoices", "Payments"):
        sheet = workbook.add_sheet(name)
        sheet.write(0, 0, "SYN-XLS-001")
        sheet.write(0, 1, 150)

    workbook.save(str(path))

    units = list(read_document(path, source_for(path), settings=options()))

    assert [
        unit.position.sheet_name for unit in units if unit.kind == "sheet"
    ] == ["Invoices", "Payments"]

    rows = [unit for unit in units if unit.kind == "row"]
    assert len(rows) == 2
    assert rows[0].cells[0].value == "SYN-XLS-001"
    assert Decimal(rows[0].cells[1].value) == Decimal("150")


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16"])
def test_csv_preserves_multiline_record_provenance(tmp_path, encoding):
    path = tmp_path / "invoice.csv"
    path.write_text(
        'invoice,description\nSYN-001,"line one\nline two"\n',
        encoding=encoding,
        newline="\n",
    )

    units = list(read_document(path, source_for(path), settings=options()))

    assert len(units) == 2
    assert units[1].position.row_number == 2
    assert units[1].position.physical_line_start == 2
    assert units[1].position.physical_line_end == 3
    assert units[1].cells[1].value == "line one\nline two"


def test_source_checksum_mismatch_is_rejected(tmp_path):
    path = tmp_path / "invoice.csv"
    path.write_text("invoice,total\nSYN-001,125\n", encoding="utf-8")
    source = source_for(path)
    path.write_text("invoice,total\nCHANGED,999\n", encoding="utf-8")

    with pytest.raises(ProcessingError) as caught:
        list(read_document(path, source, settings=options()))

    assert caught.value.code == "source_integrity_failed"


@pytest.mark.parametrize("extension", [".pdf", ".xlsx", ".xls", ".png"])
def test_corrupted_documents_fail_explicitly(tmp_path, extension):
    path = tmp_path / f"broken{extension}"
    path.write_bytes(b"not a valid document")

    with pytest.raises(ProcessingError):
        list(read_document(path, source_for(path), settings=options()))


def test_csv_column_limit_is_enforced(tmp_path):
    path = tmp_path / "wide.csv"
    path.write_text("a,b,c\n1,2,3\n", encoding="utf-8")

    with pytest.raises(ProcessingError) as caught:
        list(
            read_document(
                path,
                source_for(path),
                settings=options(max_columns=2),
            )
        )

    assert caught.value.code == "table_resource_limit"


def test_xlsx_actual_coordinates_are_checked(tmp_path):
    path = tmp_path / "wide.xlsx"
    workbook = openpyxl.Workbook()

    try:
        workbook.active["D1"] = "outside configured range"
        workbook.save(path)
    finally:
        workbook.close()

    with pytest.raises(ProcessingError) as caught:
        list(
            read_document(
                path,
                source_for(path),
                settings=options(max_columns=2),
            )
        )

    assert caught.value.code == "table_resource_limit"


def test_total_output_limit_is_enforced(tmp_path):
    path = tmp_path / "invoice.csv"
    path.write_text("invoice,total\nSYN-001,125\n", encoding="utf-8")

    with pytest.raises(ProcessingError) as caught:
        list(
            read_document(
                path,
                source_for(path),
                settings=options(max_total_characters=10),
            )
        )

    assert caught.value.code == "reading_text_limit"


def test_pdf_render_limit_prevents_large_allocation(tmp_path):
    path = tmp_path / "scan.pdf"
    save_scanned_pdf(path)

    with pytest.raises(ProcessingError) as caught:
        list(
            read_document(
                path,
                source_for(path),
                settings=options(max_pdf_render_pixels=100),
                ocr=lambda image: "SHOULD NOT BE CALLED",
            )
        )

    assert caught.value.code == "pdf_render_limit"


def test_missing_tesseract_is_an_explicit_failure(tmp_path):
    recognizer = TesseractOCR(
        options(tesseract_cmd=str(tmp_path / "missing-tesseract.exe"))
    )

    with Image.new("RGB", (100, 100), "white") as image:
        with pytest.raises(ProcessingError) as caught:
            recognizer(image)

    assert caught.value.code == "ocr_unavailable"


def test_ocr_timeout_kills_the_subprocess(monkeypatch):
    class HungProcess:
        returncode = None

        def poll(self):
            return self.returncode

        def wait(self, timeout=None):
            if self.returncode is None:
                raise subprocess.TimeoutExpired("synthetic", timeout)
            return self.returncode

        def kill(self):
            self.returncode = -9

    process = HungProcess()
    times = iter([0.0, 2.0])

    monkeypatch.setattr(
        ocr_module.subprocess,
        "Popen",
        lambda *args, **kwargs: process,
    )
    monkeypatch.setattr(ocr_module, "monotonic", lambda: next(times))

    with Image.new("RGB", (100, 100), "white") as image:
        with pytest.raises(ProcessingError) as caught:
            TesseractOCR(options(ocr_timeout_seconds=1))(image)

    assert caught.value.code == "ocr_timeout"
    assert process.returncode == -9


@pytest.mark.skipif(
    os.environ.get("RUN_TESSERACT_TESTS") != "1",
    reason="Set RUN_TESSERACT_TESTS=1 to invoke the installed Tesseract executable.",
)
@pytest.mark.parametrize("extension", [".pdf", ".jpg", ".png", ".tiff"])
def test_real_tesseract_reads_synthetic_documents(tmp_path, extension):
    path = tmp_path / f"real-ocr{extension}"

    if extension == ".pdf":
        save_scanned_pdf(path)
    else:
        image_format = {
            ".jpg": "JPEG",
            ".png": "PNG",
            ".tiff": "TIFF",
        }[extension]

        with invoice_image() as image:
            if extension == ".tiff":
                with image.copy() as second:
                    image.save(
                        path,
                        format=image_format,
                        save_all=True,
                        append_images=[second],
                    )
            else:
                image.save(path, format=image_format)

    units = list(
        read_document(
            path,
            source_for(path),
            settings=options(ocr_timeout_seconds=30),
        )
    )

    assert len(units) == (2 if extension == ".tiff" else 1)
    assert all("12345" in unit.text for unit in units)
    assert all("ocr_output_requires_review" in unit.warnings for unit in units)