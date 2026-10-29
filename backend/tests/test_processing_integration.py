from __future__ import annotations

import io
from contextlib import closing
from datetime import timedelta
from threading import Event
from uuid import UUID

import openpyxl
import pymupdf
import pytest
import xlwt
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import event, func, select

from app.core.config import BACKEND_DIR, Settings
from app.core.database import utc_now
from app.main import create_app
from app.models.extraction import (
    DocumentExtraction,
    DocumentExtractionHead,
    ExtractedUnit,
)
from app.models.ingestion import Document, DocumentStatus, JobStatus, ProcessingJob
from app.services.document_processor import DocumentProcessor, read_published_page
from app.services.document_reader import read_document
from app.services.reading_types import ReadingSettings
from app.workers.jobs import JobStore, ProcessingError
from app.workers.runtime import RunOutcome, WorkerContext, WorkerRuntime


class Clock:
    def __init__(self):
        self.value = utc_now() + timedelta(minutes=1)

    def __call__(self):
        return self.value

    def advance(self, seconds):
        self.value += timedelta(seconds=seconds)


@pytest.fixture
def environment(tmp_path):
    settings = Settings(
        _env_file=None,
        app_env="test",
        debug=False,
        database_url=(
            f"sqlite+pysqlite:///{(tmp_path / 'processing.db').as_posix()}"
        ),
        local_storage_path=tmp_path / "uploads",
        max_file_size_mb=10,
        max_upload_total_mb=20,
        job_max_attempts=3,
        job_lease_seconds=30,
        cors_origins=[],
    )

    config = Config(str(BACKEND_DIR / "alembic.ini"))
    config.attributes["settings"] = settings
    command.upgrade(config, "head")

    application = create_app(settings)

    with TestClient(application) as client:
        factory = application.state.session_factory
        clock = Clock()
        store = JobStore(
            factory,
            lease_seconds=30,
            retry_base_seconds=5,
            clock=clock,
        )
        yield settings, application, client, factory, clock, store


def payload_for(extension):
    if extension == ".csv":
        return b"invoice,total\nSYN-001,125\nSYN-002,150\n", 3, 1

    if extension == ".pdf":
        with pymupdf.open() as document:
            for number in range(1, 3):
                page = document.new_page()
                page.insert_text(
                    (40, 50),
                    f"SYNTHETIC DIGITAL INVOICE PAGE {number} TOTAL 125",
                )
            return document.tobytes(), 2, 2

    output = io.BytesIO()

    if extension == ".xlsx":
        workbook = openpyxl.Workbook()
        try:
            first = workbook.active
            first.title = "Invoices"
            first.append(["invoice", "total"])
            first.append(["SYN-001", 125])
            second = workbook.create_sheet("Payments")
            second.append(["SYN-001", 125])
            workbook.save(output)
        finally:
            workbook.close()
        return output.getvalue(), 5, 2

    workbook = xlwt.Workbook()
    for name in ("Invoices", "Payments"):
        sheet = workbook.add_sheet(name)
        sheet.write(0, 0, "SYN-001")
        sheet.write(0, 1, 125)
    workbook.save(output)
    return output.getvalue(), 4, 2


def upload(client, extension=".csv", payload=None):
    if payload is None:
        payload, _, _ = payload_for(extension)

    response = client.post(
        "/api/documents/upload",
        files=[
            (
                "files",
                (f"synthetic{extension}", payload, "application/octet-stream"),
            )
        ],
    )
    assert response.status_code == 202, response.text
    return UUID(response.json()["documents"][0]["id"])


def processor(**kwargs):
    return DocumentProcessor(
        reading_settings=ReadingSettings(_env_file=None),
        **kwargs,
    )


def runtime(environment, handler):
    settings, _, _, _, _, store = environment
    return WorkerRuntime(
        store,
        settings,
        worker_id="integration-worker",
        processor=handler,
    )


def new_context(environment, claim):
    settings, _, _, _, _, store = environment
    return WorkerContext(
        claim=claim,
        settings=settings,
        store=store,
        lease_lost=Event(),
    )


def job_for(factory, document_id):
    with factory() as session:
        return session.scalar(
            select(ProcessingJob).where(
                ProcessingJob.document_id == document_id
            )
        )


@pytest.mark.parametrize("extension", [".csv", ".pdf", ".xlsx", ".xls"])
def test_upload_worker_persists_actual_content(environment, extension):
    _, _, client, factory, _, _ = environment
    payload, expected_units, expected_parts = payload_for(extension)
    document_id = upload(client, extension, payload)

    worker = runtime(environment, processor(batch_units=2))
    assert worker.run_once() == RunOutcome.SUCCEEDED
    assert worker.run_once() == RunOutcome.IDLE

    response = client.get(f"/api/documents/{document_id}")
    assert response.status_code == 200
    assert response.json()["status"] == "completed"
    assert response.json()["progress_percent"] == 100
    assert response.json()["processed_units"] == expected_parts
    assert response.json()["total_units"] == expected_parts

    with factory() as session:
        extraction_id, units = read_published_page(
            session,
            document_id,
            limit=100,
        )
        assert extraction_id is not None
        assert len(units) == expected_units
        assert [unit.sequence for unit in units] == list(
            range(1, expected_units + 1)
        )
        assert all(unit.source.document_id == document_id for unit in units)
        assert any("125" in unit.text for unit in units)

        extraction = session.get(DocumentExtraction, extraction_id)
        assert extraction.ready_at is not None
        assert extraction.unit_count == expected_units
        assert extraction.part_count == expected_parts

        if extension == ".pdf":
            assert [unit.position.page_number for unit in units] == [1, 2]
        elif extension == ".csv":
            assert units[1].position.row_number == 2
            assert units[1].cells[0].value == "SYN-001"
        else:
            assert {
                unit.position.sheet_name for unit in units
            } == {"Invoices", "Payments"}


def test_partial_failure_is_hidden_and_retry_uses_new_attempt(environment):
    _, _, client, factory, clock, _ = environment
    document_id = upload(client)

    def failing_reader(path, source, **kwargs):
        with closing(read_document(path, source, **kwargs)) as units:
            yield next(units)
            raise ProcessingError(
                "temporary_reader_failure",
                "A synthetic temporary read failure occurred.",
                retryable=True,
            )

    assert runtime(
        environment,
        processor(reader=failing_reader, batch_units=1),
    ).run_once() == RunOutcome.RETRY_SCHEDULED

    with factory() as session:
        old_attempt = session.scalar(select(DocumentExtraction))
        old_id = old_attempt.id
        assert old_attempt.unit_count == 1
        assert old_attempt.ready_at is None
        assert read_published_page(session, document_id) == (None, [])

    clock.advance(5)

    assert runtime(
        environment,
        processor(batch_units=1),
    ).run_once() == RunOutcome.SUCCEEDED

    with factory() as session:
        new_id, units = read_published_page(session, document_id)
        assert new_id != old_id
        assert len(units) == 3
        assert session.scalar(
            select(func.count()).select_from(DocumentExtraction)
        ) == 2
        assert session.scalar(
            select(func.count()).select_from(ExtractedUnit).where(
                ExtractedUnit.extraction_id == old_id
            )
        ) == 1

    assert job_for(factory, document_id).attempt_count == 2


def test_sealed_output_is_hidden_until_job_success(environment):
    _, _, client, factory, clock, _ = environment
    document_id = upload(client)
    actual_processor = processor(batch_units=1)

    def fail_after_persistence(context):
        actual_processor(context)
        raise ProcessingError(
            "synthetic_finalize_failure",
            "A synthetic failure occurred before job completion.",
            retryable=True,
        )

    assert runtime(
        environment,
        fail_after_persistence,
    ).run_once() == RunOutcome.RETRY_SCHEDULED

    with factory() as session:
        candidate = session.scalar(select(DocumentExtraction))
        assert candidate.ready_at is not None
        assert session.get(DocumentExtractionHead, document_id) is not None
        assert read_published_page(session, document_id) == (None, [])

    clock.advance(5)
    assert runtime(environment, processor()).run_once() == RunOutcome.SUCCEEDED

    with factory() as session:
        _, units = read_published_page(session, document_id)
        assert len(units) == 3


def test_replaying_same_claim_does_not_duplicate_units(environment):
    _, _, client, factory, _, store = environment
    document_id = upload(client)
    claim = store.claim_next("manual-test-worker")
    assert claim is not None

    context = new_context(environment, claim)
    handler = processor(batch_units=1)

    first_result = handler(context)
    second_result = handler(context)

    assert first_result == second_result

    with factory() as session:
        assert session.scalar(
            select(func.count()).select_from(ExtractedUnit)
        ) == 3
        assert read_published_page(session, document_id) == (None, [])

    store.complete(claim, second_result)

    with factory() as session:
        _, units = read_published_page(session, document_id)
        assert len(units) == 3


def test_expired_attempt_cannot_write_after_replacement_claim(environment):
    _, _, client, factory, clock, store = environment
    document_id = upload(client)
    replacement = []

    def interrupted_reader(path, source, **kwargs):
        with closing(read_document(path, source, **kwargs)) as units:
            yield next(units)

            clock.advance(31)
            assert store.recover_expired() == 1
            clock.advance(5)

            new_claim = store.claim_next("replacement-worker")
            assert new_claim is not None
            replacement.append(new_claim)

            yield next(units)

    assert runtime(
        environment,
        processor(reader=interrupted_reader, batch_units=1),
    ).run_once() == RunOutcome.LEASE_LOST

    with factory() as session:
        old_attempt = session.scalar(select(DocumentExtraction))
        assert old_attempt.unit_count == 1
        assert read_published_page(session, document_id) == (None, [])

    claim = replacement[0]
    result = processor(batch_units=1)(new_context(environment, claim))
    store.complete(claim, result)

    with factory() as session:
        extraction_id, units = read_published_page(session, document_id)
        assert extraction_id == claim.lease_token
        assert len(units) == 3


def test_insert_failure_rolls_back_rows_and_manifest(environment):
    _, application, client, factory, _, _ = environment
    document_id = upload(client)
    engine = application.state.database_engine

    def fail_after_insert(
        connection,
        cursor,
        statement,
        parameters,
        execution_context,
        executemany,
    ):
        if statement.lstrip().upper().startswith("INSERT INTO EXTRACTED_UNITS"):
            raise RuntimeError("Synthetic failure after unit insertion")

    event.listen(engine, "after_cursor_execute", fail_after_insert)
    try:
        outcome = runtime(
            environment,
            processor(batch_units=2),
        ).run_once()
    finally:
        event.remove(engine, "after_cursor_execute", fail_after_insert)

    assert outcome == RunOutcome.RETRY_SCHEDULED

    with factory() as session:
        attempt = session.scalar(select(DocumentExtraction))
        assert attempt.unit_count == 0
        assert attempt.ready_at is None
        assert session.scalar(
            select(func.count()).select_from(ExtractedUnit)
        ) == 0
        assert read_published_page(session, document_id) == (None, [])

    assert job_for(factory, document_id).status == JobStatus.QUEUED


def test_empty_workbook_does_not_complete_merely_by_opening(environment):
    _, _, client, factory, _, _ = environment
    output = io.BytesIO()
    workbook = openpyxl.Workbook()
    try:
        workbook.save(output)
    finally:
        workbook.close()

    document_id = upload(client, ".xlsx", output.getvalue())

    assert runtime(
        environment,
        processor(batch_units=1),
    ).run_once() == RunOutcome.FAILED

    with factory() as session:
        document = session.get(Document, document_id)
        assert document.status == DocumentStatus.FAILED
        assert document.error_code == "no_extracted_content"
        assert read_published_page(session, document_id) == (None, [])


def test_published_pagination_is_bound_to_attempt(environment):
    _, _, client, factory, _, _ = environment
    document_id = upload(client)

    assert runtime(environment, processor()).run_once() == RunOutcome.SUCCEEDED

    with factory() as session:
        extraction_id, first = read_published_page(
            session,
            document_id,
            limit=1,
        )
        same_id, second = read_published_page(
            session,
            document_id,
            after_sequence=first[-1].sequence,
            limit=1,
            expected_extraction_id=extraction_id,
        )

        assert same_id == extraction_id
        assert first[0].sequence == 1
        assert second[0].sequence == 2


def test_deleting_document_cascades_extracted_content(environment):
    _, _, client, factory, _, _ = environment
    document_id = upload(client)

    assert runtime(environment, processor()).run_once() == RunOutcome.SUCCEEDED

    with factory() as session:
        document = session.get(Document, document_id)
        session.delete(document)
        session.commit()

    with factory() as session:
        for model in (
            DocumentExtraction,
            ExtractedUnit,
            DocumentExtractionHead,
        ):
            assert session.scalar(
                select(func.count()).select_from(model)
            ) == 0
