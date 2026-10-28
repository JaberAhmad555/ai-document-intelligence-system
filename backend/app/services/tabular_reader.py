from __future__ import annotations

import codecs
import csv
import io
import json
import re
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from xml.parsers import expat

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

from app.core.config import get_settings
from app.services.file_validation import validate_file
from app.services.reading_types import (
    CellValue,
    DocumentSource,
    ReadingSettings,
    ReadingUnit,
    SourcePosition,
)
from app.workers.jobs import ProcessingError


CELL_REFERENCE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")


def _limit(message: str) -> ProcessingError:
    return ProcessingError("table_resource_limit", message, retryable=False)


def _check_xlsx_coordinates(path: Path, settings: ReadingSettings) -> None:
    """Check actual XML coordinates before rows are materialized.

    The existing XLSX validator must run first. It rejects DTDs, invalid XML,
    suspicious archives, excessive decompression, and invalid workbook data.
    """

    def start_element(name, attributes):
        local_name = name.rsplit("}", 1)[-1]
        reference = attributes.get("r")

        if local_name == "row" and reference is not None:
            row_number = int(reference)
            if not 1 <= row_number <= settings.max_rows_per_sheet:
                raise _limit("A worksheet exceeds the configured row limit.")

        if local_name == "c" and reference is not None:
            match = CELL_REFERENCE.fullmatch(reference)
            if match is None:
                raise ProcessingError(
                    "invalid_cell_reference",
                    "The workbook contains an invalid cell reference.",
                    retryable=False,
                )

            letters, row_text = match.groups()
            column_number = 0
            for letter in letters:
                column_number = column_number * 26 + ord(letter) - ord("A") + 1

            if (
                column_number > settings.max_columns
                or int(row_text) > settings.max_rows_per_sheet
            ):
                raise _limit("A worksheet exceeds configured row or column limits.")

    with zipfile.ZipFile(path) as archive:
        for entry in archive.infolist():
            if entry.is_dir() or not entry.filename.lower().endswith(".xml"):
                continue

            parser = expat.ParserCreate(namespace_separator="}")
            parser.StartElementHandler = start_element

            with archive.open(entry) as member:
                while chunk := member.read(64 * 1024):
                    parser.Parse(chunk, False)
                parser.Parse(b"", True)


def _cell(
    value,
    column: int,
    row: int,
    settings: ReadingSettings,
    *,
    formula: bool = False,
    error: bool = False,
) -> CellValue:
    address = f"{get_column_letter(column)}{row}"
    expression = None

    if formula:
        data_type = "formula"
        expression = str(value)
        result = None
    elif error:
        data_type = "error"
        result = str(value)
    elif value is None:
        data_type = "blank"
        result = None
    elif isinstance(value, bool):
        data_type = "boolean"
        result = "true" if value else "false"
    elif isinstance(value, (datetime, date, time)):
        data_type = "date"
        result = value.isoformat()
    elif isinstance(value, (int, float, Decimal)):
        number = Decimal(str(value))
        if not number.is_finite():
            raise ProcessingError(
                "invalid_numeric_cell",
                "A spreadsheet contains a non-finite numeric value.",
                retryable=False,
            )
        data_type = "number"
        result = format(number, "f")
    else:
        data_type = "text"
        result = str(value)

    for text in (result, expression):
        if text is not None and len(text) > settings.max_cell_characters:
            raise _limit("A table cell exceeds the configured character limit.")

    return CellValue(
        column_number=column,
        address=address,
        data_type=data_type,
        value=result,
        formula=expression,
    )


def _row_unit(
    source: DocumentSource,
    position: SourcePosition,
    cells: tuple[CellValue, ...],
    *,
    method: str,
) -> ReadingUnit:
    flags = []
    if any(cell.data_type == "formula" for cell in cells):
        flags.append("formula_value_unavailable")
    if any(cell.data_type == "error" for cell in cells):
        flags.append("spreadsheet_error_cell")

    values = [
        cell.formula if cell.data_type == "formula" else cell.value
        for cell in cells
    ]

    return ReadingUnit(
        source=source,
        position=position,
        kind="row",
        method=method,
        text=json.dumps(values, ensure_ascii=False),
        cells=cells,
        warnings=tuple(flags),
    )


def read_xlsx(
    path: Path,
    source: DocumentSource,
    settings: ReadingSettings,
) -> Iterator[ReadingUnit]:
    validate_file(path, ".xlsx", get_settings())
    _check_xlsx_coordinates(path, settings)

    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )

    try:
        sheets = workbook.worksheets
        if not 1 <= len(sheets) <= settings.max_pages:
            raise _limit("The workbook sheet count exceeds the configured limit.")

        for sheet_index, sheet in enumerate(sheets, start=1):
            yield ReadingUnit(
                source=source,
                position=SourcePosition(
                    part_number=sheet_index,
                    part_count=len(sheets),
                    sheet_name=sheet.title,
                ),
                kind="sheet",
                method="spreadsheet",
                text="",
            )

            # Actual coordinates were checked above. Ignore inaccurate
            # worksheet dimension metadata when streaming the rows.
            sheet.reset_dimensions()

            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                if row_number > settings.max_rows_per_sheet:
                    raise _limit("A worksheet exceeds the configured row limit.")
                if len(row) > settings.max_columns:
                    raise _limit("A worksheet exceeds the configured column limit.")

                cells = tuple(
                    _cell(
                        cell.value,
                        column,
                        row_number,
                        settings,
                        formula=cell.data_type == "f",
                        error=cell.data_type == "e",
                    )
                    for column, cell in enumerate(row, start=1)
                )

                yield _row_unit(
                    source,
                    SourcePosition(
                        part_number=sheet_index,
                        part_count=len(sheets),
                        sheet_name=sheet.title,
                        row_number=row_number,
                    ),
                    cells,
                    method="spreadsheet",
                )
    finally:
        workbook.close()


def read_xls(
    path: Path,
    source: DocumentSource,
    settings: ReadingSettings,
) -> Iterator[ReadingUnit]:
    workbook = xlrd.open_workbook(
        filename=str(path),
        on_demand=True,
        logfile=io.StringIO(),
    )

    try:
        if not 1 <= workbook.nsheets <= settings.max_pages:
            raise _limit("The workbook sheet count exceeds the configured limit.")

        for sheet_index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(sheet_index)

            try:
                if (
                    sheet.nrows > settings.max_rows_per_sheet
                    or sheet.ncols > settings.max_columns
                ):
                    raise _limit("A worksheet exceeds configured table limits.")

                yield ReadingUnit(
                    source=source,
                    position=SourcePosition(
                        part_number=sheet_index + 1,
                        part_count=workbook.nsheets,
                        sheet_name=sheet.name,
                    ),
                    kind="sheet",
                    method="spreadsheet",
                    text="",
                    warnings=("xls_formula_expressions_unavailable",),
                )

                for row_index in range(sheet.nrows):
                    cells = []

                    for column_index in range(sheet.ncols):
                        cell = sheet.cell(row_index, column_index)
                        value = cell.value
                        is_error = cell.ctype == xlrd.XL_CELL_ERROR

                        if cell.ctype in {
                            xlrd.XL_CELL_EMPTY,
                            xlrd.XL_CELL_BLANK,
                        }:
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            converted = xlrd.xldate_as_datetime(
                                value,
                                workbook.datemode,
                            )
                            value = (
                                converted.time()
                                if 0 <= cell.value < 1
                                else converted
                            )
                        elif is_error:
                            value = xlrd.error_text_from_code.get(
                                value,
                                "Unknown spreadsheet error",
                            )

                        cells.append(
                            _cell(
                                value,
                                column_index + 1,
                                row_index + 1,
                                settings,
                                error=is_error,
                            )
                        )

                    yield _row_unit(
                        source,
                        SourcePosition(
                            part_number=sheet_index + 1,
                            part_count=workbook.nsheets,
                            sheet_name=sheet.name,
                            row_number=row_index + 1,
                        ),
                        tuple(cells),
                        method="spreadsheet",
                    )
            finally:
                workbook.unload_sheet(sheet_index)
    finally:
        workbook.release_resources()


def read_csv(
    path: Path,
    source: DocumentSource,
    settings: ReadingSettings,
) -> Iterator[ReadingUnit]:
    with path.open("rb") as stream:
        prefix = stream.read(4)

    if prefix.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        raise ProcessingError(
            "unsupported_csv_encoding",
            "CSV reading supports UTF-8 and BOM-marked UTF-16.",
            retryable=False,
        )

    encoding = (
        "utf-16"
        if prefix.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE))
        else "utf-8-sig"
    )

    with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
        sample = stream.read(8192)
        stream.seek(0)

        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(stream, dialect=dialect, strict=True)
        expected_columns = None
        previous_line = 0

        for row_number, values in enumerate(reader, start=1):
            line_start = previous_line + 1
            previous_line = reader.line_num

            if row_number > settings.max_rows_per_sheet:
                raise _limit("The CSV exceeds the configured row limit.")
            if len(values) > settings.max_columns:
                raise _limit("The CSV exceeds the configured column limit.")

            if values and any(value.strip() for value in values):
                if expected_columns is None:
                    expected_columns = len(values)
                elif len(values) != expected_columns:
                    raise ProcessingError(
                        "inconsistent_csv_columns",
                        "CSV records contain inconsistent column counts.",
                        retryable=False,
                    )

            for value in values:
                if any(
                    (ord(character) < 32 and character not in "\t\r\n")
                    or ord(character) == 127
                    for character in value
                ):
                    raise ProcessingError(
                        "binary_csv",
                        "The CSV contains unsupported control characters.",
                        retryable=False,
                    )

            cells = tuple(
                _cell(value, column, row_number, settings)
                for column, value in enumerate(values, start=1)
            )

            yield _row_unit(
                source,
                SourcePosition(
                    part_number=1,
                    part_count=1,
                    row_number=row_number,
                    physical_line_start=line_start,
                    physical_line_end=reader.line_num,
                ),
                cells,
                method="csv",
            )